"""
Load historical results + bookmaker odds from tennis-data.co.uk into the BigQuery bronze layer.

Purpose: BENCHMARKING ONLY. These odds are never model inputs. They are used to compare the
model with the market (log loss, calibration) and to backtest betting rules out-of-sample.

Coverage: ATP and WTA main tour (no Challengers / ITF), one Excel file per tour per year.
Bookmaker columns vary by year; the useful ones for us are:
  PSW / PSL     Pinnacle closing odds (sharpest benchmark available)
  MaxW / MaxL   best price across bookmakers
  AvgW / AvgL   average price across bookmakers
  B365W / B365L Bet365
  BFEW / BFEL   Betfair Exchange (later years only)

Bronze rules (same as load_tml.py): every column stored as STRING, append-only with lineage
columns, unchanged files skipped using bronze._ingestion_log.

Usage (from the repo root, venv active):
  python ingestion/load_tennis_data.py --dry-run                 # download + parse only
  python ingestion/load_tennis_data.py                           # 2013 -> this year, ATP + WTA
  python ingestion/load_tennis_data.py --first-year 2024 --tours atp
  python ingestion/load_tennis_data.py --local-dir C:\\Users\\colmh\\Downloads\\tennis_data
      (if the site refuses scripted downloads: save the files from your browser as
       atp_2025.xlsx, wta_2025.xlsx, ... into that folder and load them from there)
"""

from __future__ import annotations

import argparse
import hashlib
import io
import logging
import re
import sys
import time
from datetime import date, datetime, timezone
from pathlib import Path

import requests

from load_tml import (DEFAULT_DATASET, DEFAULT_LOCATION, DEFAULT_PROJECT, HTTP_TIMEOUT_SECONDS,
                      USER_AGENT, BronzeWriter)

BASE_URL = "https://www.tennis-data.co.uk"
DEFAULT_FIRST_YEAR = 2013
REQUEST_PAUSE_SECONDS = 2.0      # small free site: be gentle

TOURS = {
    # tour: (bronze table, path template relative to BASE_URL, without extension)
    "atp": ("tennis_data_atp_matches", "{year}/{year}"),
    "wta": ("tennis_data_wta_matches", "{year}w/{year}"),
}

log = logging.getLogger("load_tennis_data")


# --------------------------------------------------------------------------------------
# Download
# --------------------------------------------------------------------------------------

def fetch_remote(session: requests.Session, stem: str) -> tuple[str, bytes] | None:
    """Try <stem>.xlsx then <stem>.xls. Returns (path, content) or None if neither exists."""
    for ext in (".xlsx", ".xls"):
        path = stem + ext
        resp = session.get(f"{BASE_URL}/{path}", timeout=HTTP_TIMEOUT_SECONDS)
        time.sleep(REQUEST_PAUSE_SECONDS)
        if resp.status_code == 404:
            continue
        if resp.status_code == 403:
            raise PermissionError(
                "403 Forbidden: the site refused a scripted download. Download the files in "
                "your browser and use --local-dir instead.")
        resp.raise_for_status()
        return path, resp.content
    return None


def fetch_local(local_dir: Path, tour: str, year: int, stem: str) -> tuple[str, bytes] | None:
    """Files saved by hand as <tour>_<year>.xlsx / .xls. Logged under the site path so the
    unchanged-file check works the same whichever way a file was obtained."""
    for ext in (".xlsx", ".xls"):
        f = local_dir / f"{tour}_{year}{ext}"
        if f.exists():
            return stem + ext, f.read_bytes()
    return None


# --------------------------------------------------------------------------------------
# Parse
# --------------------------------------------------------------------------------------

def clean_column(name: object, used: set[str]) -> str:
    """BigQuery-safe column name: 'Best of' -> 'Best_of', 'B&WW' -> 'B_WW'."""
    col = re.sub(r"[^0-9A-Za-z_]", "_", str(name).strip()) or "col"
    if col[0].isdigit():
        col = "_" + col
    base, n = col, 2
    while col.lower() in used:
        col, n = f"{base}_{n}", n + 1
    used.add(col.lower())
    return col


def to_text(value: object) -> str | None:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    text = str(value).strip()
    return text or None


def read_rows(path: str, content: bytes) -> list[list[object]]:
    if path.endswith(".xlsx"):
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        return [list(r) for r in wb.worksheets[0].iter_rows(values_only=True)]
    import xlrd   # only needed for the older .xls files
    book = xlrd.open_workbook(file_contents=content)
    sheet = book.sheet_by_index(0)
    rows = []
    for i in range(sheet.nrows):
        row = []
        for cell in sheet.row(i):
            if cell.ctype == xlrd.XL_CELL_DATE:
                row.append(xlrd.xldate_as_datetime(cell.value, book.datemode))
            elif cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                row.append(None)
            else:
                row.append(cell.value)
        rows.append(row)
    return rows


def parse_workbook(path: str, content: bytes) -> tuple[list[str], list[dict]]:
    raw = read_rows(path, content)
    if not raw:
        return [], []
    used: set[str] = set()
    header = [clean_column(h, used) if h is not None else None for h in raw[0]]
    columns = [h for h in header if h]
    rows = []
    for values in raw[1:]:
        row = {h: to_text(v) for h, v in zip(header, values) if h}
        if any(v is not None for v in row.values()):     # skip blank trailing rows
            rows.append(row)
    return columns, rows


def add_lineage(rows: list[dict], path: str, md5: str, loaded_at: str) -> list[dict]:
    for i, row in enumerate(rows, start=1):
        row.update({"_source_file": f"tennis_data/{path}", "_source_url": f"{BASE_URL}/{path}",
                    "_source_row_number": i, "_source_md5": md5, "_loaded_at": loaded_at})
    return rows


# --------------------------------------------------------------------------------------
# Main
# --------------------------------------------------------------------------------------

def run(args: argparse.Namespace) -> int:
    years = range(args.first_year, (args.last_year or date.today().year) + 1)
    writer = None if args.dry_run else BronzeWriter(args.project, args.dataset, args.location)
    seen_md5 = {} if (writer is None or args.force) else writer.latest_md5s()
    local_dir = Path(args.local_dir) if args.local_dir else None

    session = requests.Session()
    session.headers["User-Agent"] = USER_AGENT

    loaded = unchanged = missing = total_rows = 0
    failures: list[str] = []

    for tour in args.tours:
        table, template = TOURS[tour]
        log.info("%s -> bronze.%s", tour, table)
        for year in years:
            stem = template.format(year=year)
            try:
                got = (fetch_local(local_dir, tour, year, stem) if local_dir
                       else fetch_remote(session, stem))
                if got is None:
                    missing += 1
                    log.info("  %s %d: not found, skipping", tour, year)
                    continue
                path, content = got
                log_key = f"tennis_data/{path}"
                md5 = hashlib.md5(content).hexdigest()
                if seen_md5.get(log_key) == md5:
                    unchanged += 1
                    log.info("  %s unchanged, skipping", path)
                    continue

                columns, rows = parse_workbook(path, content)
                if not rows:
                    log.info("  %s is empty, skipping", path)
                    continue
                loaded_at = datetime.now(timezone.utc).isoformat()
                rows = add_lineage(rows, path, md5, loaded_at)

                if writer:
                    writer.load(table, columns, rows, "append")
                    writer.record(log_key, table, md5, len(rows), loaded_at)
                loaded += 1
                total_rows += len(rows)
                # Column names and counts only: never print odds values (public run logs).
                odds_cols = [c for c in ("PSW", "MaxW", "AvgW", "B365W", "BFEW") if c in columns]
                log.info("  %s: %d rows, %d columns, odds columns %s%s", path, len(rows),
                         len(columns), odds_cols, " (dry run)" if args.dry_run else "")
            except PermissionError as exc:
                log.error("  %s %d: %s", tour, year, exc)
                return 1
            except Exception as exc:   # keep going; report at the end
                failures.append(f"{tour} {year}")
                log.error("  %s %d FAILED: %s", tour, year, exc)

    log.info("Done. files loaded=%d unchanged=%d not_found=%d rows=%d failed=%d",
             loaded, unchanged, missing, total_rows, len(failures))
    if failures:
        log.error("Failed: %s", ", ".join(failures))
        return 1
    return 0


def parse_args(argv: list[str] | None = None) -> argparse.Namespace:
    p = argparse.ArgumentParser(description="Load tennis-data.co.uk results + odds into bronze.")
    p.add_argument("--tours", nargs="+", choices=list(TOURS), default=list(TOURS))
    p.add_argument("--first-year", type=int, default=DEFAULT_FIRST_YEAR)
    p.add_argument("--last-year", type=int, default=None, help="default: this year")
    p.add_argument("--local-dir", default=None,
                   help="load <tour>_<year>.xlsx files saved by hand instead of downloading")
    p.add_argument("--project", default=DEFAULT_PROJECT)
    p.add_argument("--dataset", default=DEFAULT_DATASET)
    p.add_argument("--location", default=DEFAULT_LOCATION)
    p.add_argument("--dry-run", action="store_true", help="download and parse only; no BigQuery")
    p.add_argument("--force", action="store_true", help="reload files even if unchanged")
    return p.parse_args(argv)


if __name__ == "__main__":
    args = parse_args()
    logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)s %(message)s")
    sys.exit(run(args))